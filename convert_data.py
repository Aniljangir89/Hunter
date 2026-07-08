"""
Prepare the HR dataset: convert Excel → enhanced JSON with IDs, 
email type classification, and validation placeholders.
"""
import openpyxl
import json
import os
import re
from datetime import datetime

PERSONAL_DOMAINS = {
    'gmail.com', 'yahoo.com', 'yahoo.in', 'yahoo.co.in',
    'hotmail.com', 'outlook.com', 'live.com', 'msn.com',
    'rediffmail.com', 'aol.com', 'icloud.com', 'protonmail.com',
    'mail.com', 'zoho.com', 'ymail.com', 'googlemail.com'
}

# City name normalization mapping
CITY_NORMALIZE = {
    'bengaluru': 'Bangalore',
    'banglore': 'Bangalore',
    'bangluru': 'Bangalore',
    'bangaluru': 'Bangalore',
    'banglaore': 'Bangalore',
    'blr': 'Bangalore',
    'gurgaon': 'Gurgaon',
    'gurugram': 'Gurgaon',
    'gurugam': 'Gurgaon',
    'noida': 'Noida',
    'greater noida': 'Noida',
    'mumbai': 'Mumbai',
    'bombay': 'Mumbai',
    'pune': 'Pune',
    'delhi': 'Delhi',
    'new delhi': 'Delhi',
    'delhi ncr': 'Delhi NCR',
    'ncr': 'Delhi NCR',
    'hyderabad': 'Hyderabad',
    'chennai': 'Chennai',
    'madras': 'Chennai',
    'kolkata': 'Kolkata',
    'calcutta': 'Kolkata',
    'jaipur': 'Jaipur',
    'ahmedabad': 'Ahmedabad',
    'chandigarh': 'Chandigarh',
    'lucknow': 'Lucknow',
    'indore': 'Indore',
    'kochi': 'Kochi',
    'cochin': 'Kochi',
    'thiruvananthapuram': 'Thiruvananthapuram',
    'trivandrum': 'Thiruvananthapuram',
    'coimbatore': 'Coimbatore',
    'vadodara': 'Vadodara',
    'baroda': 'Vadodara',
    'remote': 'Remote',
    'work from home': 'Remote',
    'wfh': 'Remote',
    'pan india': 'Pan India',
    'india': 'Pan India',
    'across india': 'Pan India',
}

def normalize_city(city):
    """Normalize a city name using the mapping."""
    city_lower = city.strip().lower()
    return CITY_NORMALIZE.get(city_lower, city.strip().title())

def classify_email(email, domain):
    """Classify email as corporate or personal."""
    if not domain:
        return 'invalid'
    if domain.lower() in PERSONAL_DOMAINS:
        return 'personal'
    return 'corporate'

def is_valid_email_syntax(email):
    """Basic email syntax validation."""
    pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    return bool(re.match(pattern, str(email)))

def main():
    data_dir = "Data"
    files = os.listdir(data_dir)
    filepath = os.path.join(data_dir, files[0])
    
    print(f"Loading: {filepath}")
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    
    records = []
    now = datetime.now().isoformat()
    
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        company, location, email = row
        
        if not company and not email:
            continue
        
        company = str(company).strip() if company else 'Unknown'
        email = str(email).strip().lower() if email else ''
        location = str(location).strip() if location else 'Unknown'
        
        # Extract domain
        domain = email.split('@')[-1] if '@' in email else ''
        
        # Parse and normalize locations
        raw_locations = [loc.strip() for loc in location.replace(',', '/').split('/') if loc.strip()]
        if not raw_locations:
            raw_locations = ['Unknown']
        normalized_locations = [normalize_city(loc) for loc in raw_locations]
        
        # Classify email
        email_type = classify_email(email, domain)
        syntax_valid = is_valid_email_syntax(email)
        
        records.append({
            'id': idx,
            'company': company,
            'location': location,
            'locations': normalized_locations,
            'email': email,
            'domain': domain,
            'email_type': email_type,
            'validation': {
                'syntax': syntax_valid,
                'mx': None,
                'smtp': None,
                'status': 'valid_syntax' if syntax_valid else 'invalid_syntax'
            },
            'added_at': now,
            'is_cleaned': False
        })
    
    # Stats
    print(f"\nTotal records: {len(records)}")
    
    types = {}
    for r in records:
        t = r['email_type']
        types[t] = types.get(t, 0) + 1
    print(f"Email types: {types}")
    
    syntax_valid = sum(1 for r in records if r['validation']['syntax'])
    print(f"Valid syntax: {syntax_valid}/{len(records)}")
    
    cities = {}
    for r in records:
        for loc in r['locations']:
            cities[loc] = cities.get(loc, 0) + 1
    top_cities = sorted(cities.items(), key=lambda x: -x[1])[:15]
    print(f"Top 15 cities (normalized): {top_cities}")
    
    # Find duplicates
    emails_seen = {}
    dupes = 0
    for r in records:
        if r['email'] in emails_seen:
            dupes += 1
        else:
            emails_seen[r['email']] = r['id']
    print(f"Duplicate emails: {dupes}")
    
    with open('data.json', 'w', encoding='utf-8') as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
    
    print(f"\nSaved {len(records)} records to data.json successfully!")

if __name__ == '__main__':
    main()
